from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook

from financial_model_engine.domain import ModelAssumptions
from financial_model_engine.engine import ModelResult
from financial_model_engine.excel.cell_mapper import CellMapper
from financial_model_engine.excel.sheets import (
    write_assumptions_sheet,
    write_checks_sheet,
    write_cover_sheet,
    write_line_items_sheet,
)


@dataclass(frozen=True)
class WorkbookCompileResult:
    output_path: Path
    cell_mapper: CellMapper


def compile_workbook(
    assumptions: ModelAssumptions, result: ModelResult, output_path: Path
) -> WorkbookCompileResult:
    """Render an already-computed ModelResult into an .xlsx workbook.

    Reads values from `assumptions` and `result`; performs no financial
    calculations of its own.
    """
    cell_mapper = CellMapper()

    workbook = Workbook()
    cover_sheet = workbook.active
    cover_sheet.title = "Cover"
    assumptions_sheet = workbook.create_sheet("Assumptions")
    operating_sheet = workbook.create_sheet("Operating Model")
    income_statement_sheet = workbook.create_sheet("Income Statement")
    balance_sheet_sheet = workbook.create_sheet("Balance Sheet")
    cash_flow_sheet = workbook.create_sheet("Cash Flow")
    checks_sheet = workbook.create_sheet("Model Checks")

    write_cover_sheet(cover_sheet, result)
    write_assumptions_sheet(assumptions_sheet, assumptions, cell_mapper)
    write_line_items_sheet(
        operating_sheet,
        "Operating Model",
        "Operating Model",
        result.periods,
        result.operating_schedule,
        cell_mapper,
    )
    write_line_items_sheet(
        income_statement_sheet,
        "Income Statement",
        "Income Statement",
        result.periods,
        result.income_statement.line_items,
        cell_mapper,
    )
    write_line_items_sheet(
        balance_sheet_sheet,
        "Balance Sheet",
        "Balance Sheet",
        result.periods,
        result.balance_sheet.line_items,
        cell_mapper,
    )
    write_line_items_sheet(
        cash_flow_sheet,
        "Cash Flow",
        "Cash Flow",
        result.periods,
        result.cash_flow.line_items,
        cell_mapper,
    )
    write_checks_sheet(checks_sheet, result.checks)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)

    return WorkbookCompileResult(output_path=output_path, cell_mapper=cell_mapper)
