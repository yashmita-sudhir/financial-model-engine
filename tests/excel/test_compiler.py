from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from financial_model_engine.domain import FinancialStatementType, ModelAssumptions
from financial_model_engine.domain import ForecastPeriod as Period
from financial_model_engine.engine import (
    LineItemResult,
    ModelCheck,
    ModelResult,
    PeriodValue,
    StatementResult,
)
from financial_model_engine.excel import SHEET_NAMES, compile_workbook


def make_assumptions() -> ModelAssumptions:
    return ModelAssumptions(
        starting_customers=Decimal("10"),
        monthly_price=Decimal("100"),
        customer_growth=Decimal("0.15"),
        monthly_marketing=Decimal("5000"),
        monthly_payroll=Decimal("8000"),
        other_opex=Decimal("1000"),
        initial_equity=Decimal("50000"),
        tax_rate=Decimal("0.20"),
        forecast_months=12,
    )


def pv(period_id: str, value: str) -> PeriodValue:
    return PeriodValue(period_id=period_id, value=Decimal(value))


def make_result() -> ModelResult:
    periods = [
        Period(id="2027-01", index=0, label="Jan 2027", year=2027, month=1),
        Period(id="2027-02", index=1, label="Feb 2027", year=2027, month=2),
    ]
    operating_schedule = [
        LineItemResult(
            line_item_id="customers",
            values=[pv("2027-01", "10"), pv("2027-02", "11.5")],
        )
    ]
    income_statement = StatementResult(
        statement_type=FinancialStatementType.INCOME_STATEMENT,
        line_items=[
            LineItemResult(
                line_item_id="revenue",
                values=[pv("2027-01", "1000"), pv("2027-02", "1150")],
            )
        ],
    )
    balance_sheet = StatementResult(
        statement_type=FinancialStatementType.BALANCE_SHEET,
        line_items=[
            LineItemResult(
                line_item_id="cash",
                values=[pv("2027-01", "46000"), pv("2027-02", "42150")],
            )
        ],
    )
    cash_flow = StatementResult(
        statement_type=FinancialStatementType.CASH_FLOW,
        line_items=[
            LineItemResult(
                line_item_id="ending_cash",
                values=[pv("2027-01", "46000"), pv("2027-02", "42150")],
            )
        ],
    )
    checks = [
        ModelCheck(
            id="balance_sheet_balances",
            name="Balance Sheet Balances",
            passed=True,
            difference=Decimal("0"),
        )
    ]
    return ModelResult(
        periods=periods,
        operating_schedule=operating_schedule,
        income_statement=income_statement,
        balance_sheet=balance_sheet,
        cash_flow=cash_flow,
        checks=checks,
    )


def test_workbook_compiles_successfully(tmp_path: Path):
    output = compile_workbook(make_assumptions(), make_result(), tmp_path / "model.xlsx")
    assert output.output_path.exists()


def test_sheet_names_and_order(tmp_path: Path):
    compile_workbook(make_assumptions(), make_result(), tmp_path / "model.xlsx")
    wb = openpyxl.load_workbook(tmp_path / "model.xlsx")
    assert wb.sheetnames == SHEET_NAMES


def test_assumptions_written_correctly(tmp_path: Path):
    output = compile_workbook(make_assumptions(), make_result(), tmp_path / "model.xlsx")
    wb = openpyxl.load_workbook(tmp_path / "model.xlsx")
    ws = wb["Assumptions"]
    sheet_name, cell_ref = output.cell_mapper.get_assumption_cell("monthly_price")
    assert sheet_name == "Assumptions"
    assert ws[cell_ref].value == 100.0


def test_period_headers_correct(tmp_path: Path):
    compile_workbook(make_assumptions(), make_result(), tmp_path / "model.xlsx")
    wb = openpyxl.load_workbook(tmp_path / "model.xlsx")
    ws = wb["Operating Model"]
    assert ws["B3"].value == "Jan 2027"
    assert ws["C3"].value == "Feb 2027"


def test_operating_schedule_values_written(tmp_path: Path):
    output = compile_workbook(make_assumptions(), make_result(), tmp_path / "model.xlsx")
    wb = openpyxl.load_workbook(tmp_path / "model.xlsx")
    ws = wb["Operating Model"]
    cell_ref = output.cell_mapper.get_line_item_cell(
        "customers", "2027-02", "Operating Model"
    )
    assert ws[cell_ref].value == 11.5


def test_income_statement_values_written(tmp_path: Path):
    output = compile_workbook(make_assumptions(), make_result(), tmp_path / "model.xlsx")
    wb = openpyxl.load_workbook(tmp_path / "model.xlsx")
    ws = wb["Income Statement"]
    cell_ref = output.cell_mapper.get_line_item_cell(
        "revenue", "2027-01", "Income Statement"
    )
    assert ws[cell_ref].value == 1000.0


def test_balance_sheet_values_written(tmp_path: Path):
    output = compile_workbook(make_assumptions(), make_result(), tmp_path / "model.xlsx")
    wb = openpyxl.load_workbook(tmp_path / "model.xlsx")
    ws = wb["Balance Sheet"]
    cell_ref = output.cell_mapper.get_line_item_cell(
        "cash", "2027-02", "Balance Sheet"
    )
    assert ws[cell_ref].value == 42150.0


def test_cash_flow_values_written(tmp_path: Path):
    output = compile_workbook(make_assumptions(), make_result(), tmp_path / "model.xlsx")
    wb = openpyxl.load_workbook(tmp_path / "model.xlsx")
    ws = wb["Cash Flow"]
    cell_ref = output.cell_mapper.get_line_item_cell(
        "ending_cash", "2027-01", "Cash Flow"
    )
    assert ws[cell_ref].value == 46000.0


def test_model_checks_written(tmp_path: Path):
    compile_workbook(make_assumptions(), make_result(), tmp_path / "model.xlsx")
    wb = openpyxl.load_workbook(tmp_path / "model.xlsx")
    ws = wb["Model Checks"]
    assert ws["A4"].value == "Balance Sheet Balances"
    assert ws["C4"].value == 0.0
    assert ws["D4"].value == "-"


def test_pass_status_rendered():
    from financial_model_engine.excel.sheets import write_checks_sheet
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    write_checks_sheet(
        ws,
        [
            ModelCheck(id="c1", name="Passing Check", passed=True),
            ModelCheck(id="c2", name="Failing Check", passed=False, message="off"),
        ],
    )
    assert ws["B4"].value == "PASS"
    assert ws["B5"].value == "FAIL"
    assert ws["D5"].value == "off"


def test_line_item_cell_mapper_roundtrip():
    from financial_model_engine.excel import CellMapper

    mapper = CellMapper()
    mapper.map_line_item("revenue", "2027-01", "Income Statement", "B7")
    assert mapper.get_line_item_cell("revenue", "2027-01", "Income Statement") == "B7"


def test_assumption_cell_mapper_roundtrip():
    from financial_model_engine.excel import CellMapper

    mapper = CellMapper()
    mapper.map_assumption("monthly_price", "Assumptions", "B3")
    assert mapper.get_assumption_cell("monthly_price") == ("Assumptions", "B3")


def test_decimal_values_written_safely(tmp_path: Path):
    output = compile_workbook(make_assumptions(), make_result(), tmp_path / "model.xlsx")
    wb = openpyxl.load_workbook(tmp_path / "model.xlsx")
    ws = wb["Operating Model"]
    cell_ref = output.cell_mapper.get_line_item_cell(
        "customers", "2027-01", "Operating Model"
    )
    value = ws[cell_ref].value
    assert isinstance(value, (int, float))
    assert value == 10


def test_compiler_does_not_mutate_model_result(tmp_path: Path):
    result = make_result()
    before = result.model_dump()
    compile_workbook(make_assumptions(), result, tmp_path / "model.xlsx")
    assert result.model_dump() == before


def test_compiler_does_not_mutate_assumptions(tmp_path: Path):
    assumptions = make_assumptions()
    before = assumptions.model_dump()
    compile_workbook(assumptions, make_result(), tmp_path / "model.xlsx")
    assert assumptions.model_dump() == before


def test_output_written_under_requested_directory(tmp_path: Path):
    nested = tmp_path / "nested" / "dir" / "model.xlsx"
    output = compile_workbook(make_assumptions(), make_result(), nested)
    assert output.output_path == nested
    assert nested.exists()
